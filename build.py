#!/usr/bin/env python3
"""
新荏田西ポータル — Excelからindex.htmlのリンクデータを更新するスクリプト

Excel列構成（1行目はヘッダー）:
  A列: カテゴリ名
  B列: サービス名
  C列: URL
  D列: 説明（省略可）
  E列: アイコン画像（セルに画像を挿入 → 自動抽出して icons/<ドメイン>.png に保存）

アイコンは URL のドメインで決まる（例: https://loilonote.app/... → icons/loilonote.app.png）。
行の追加・並べ替えをしてもアイコンはずれない。同じドメインの画像がすでにあればそれを使う。

使い方:
  python3 build.py [Excelファイルパス]
  例: python3 build.py data.xlsx
"""

import sys, os, re, zipfile, shutil, json, tempfile
from pathlib import Path
from html import escape

try:
    from openpyxl import load_workbook
except ImportError:
    print("エラー: openpyxl が必要です。以下を実行してください:")
    print("  pip3 install openpyxl")
    sys.exit(1)

import xml.etree.ElementTree as ET

# ────────── 設定 ──────────
EXCEL_PATH    = sys.argv[1] if len(sys.argv) > 1 else "data.xlsx"
INDEX_HTML    = Path(__file__).parent / "index.html"
ICONS_DIR     = Path(__file__).parent / "icons"
EMBED_DIR     = Path(tempfile.mkdtemp())   # Excel から取り出した画像の一時置き場（行番号で管理）
DATA_START_ROW = 2

# カテゴリ名 → Lucide アイコン名・カラークラス
CATEGORY_META = {
    "学習管理":           {"icon": "book-open",      "color": "tg-blue"},
    "児童管理":           {"icon": "users",           "color": "tg-green"},
    "デジタル教科書":     {"icon": "monitor",         "color": "tg-purple"},
    "校務管理":           {"icon": "folder-open",     "color": "tg-slate"},
    "生成AI":             {"icon": "sparkles",        "color": "tg-amber"},
    "研修管理":           {"icon": "clipboard-list",  "color": "tg-teal"},
    "授業準備・教材作成": {"icon": "pencil-ruler",    "color": "tg-rose"},
    "便利ツール":         {"icon": "wrench",          "color": "tg-indigo"},
}

FALLBACK_ICONS  = ["layout-grid", "bookmark", "star", "tag", "lightbulb", "target", "zap", "globe"]
FALLBACK_COLORS = ["tg-blue", "tg-green", "tg-purple", "tg-slate", "tg-amber", "tg-teal", "tg-rose", "tg-indigo"]


# ────────── URL からドメインを抽出 ──────────
def extract_domain(url: str) -> str:
    if not url or url == "#":
        return ""
    m = re.match(r"https?://([^/]+)", url)
    if not m:
        return ""
    host = m.group(1)
    return host[4:] if host.startswith("www.") else host


# ────────── 埋め込み画像を抽出 ──────────
def _resolve_sheet_xml(z, sheet_name: str) -> str:
    """シート名から xl/worksheets/sheetN.xml のパスを返す。見つからなければ sheet1.xml。"""
    try:
        wb_root = ET.fromstring(z.read("xl/workbook.xml"))
        rels_root = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        rid_to_target = {rel.get("Id"): rel.get("Target") for rel in rels_root}
        for elem in wb_root.iter():
            if elem.tag.split("}")[-1] == "sheet" and elem.get("name") == sheet_name:
                rid = elem.get(f"{{{NS_R}}}id")
                target = rid_to_target.get(rid, "")
                if target:
                    return f"xl/{target}"
    except Exception:
        pass
    return "xl/worksheets/sheet1.xml"


def extract_embedded_images(xlsx_path: str, sheet_name: str = "リンク一覧") -> dict:
    row_to_image = {}
    EMBED_DIR.mkdir(exist_ok=True)
    try:
        with zipfile.ZipFile(xlsx_path, "r") as z:
            names = z.namelist()
            media_files = [n for n in names if re.match(r"xl/media/image\d+\.", n)]
            if not media_files:
                return row_to_image
            print(f"  xlsx 内の画像ファイル数: {len(media_files)}")

            sheet_xml_path = _resolve_sheet_xml(z, sheet_name)
            print(f"  対象シートXML: {sheet_xml_path}")

            # Excel 365「セルに配置」形式
            if "xl/richData/richValueRel.xml" in names and "xl/metadata.xml" in names:
                result = _extract_rich_value_images(z, names, sheet_xml_path)
                if result:
                    return result

            # 旧来の floating image 形式
            return _extract_drawing_images(z, names)
    except Exception as e:
        print(f"  ⚠ 画像抽出エラー: {e}")
    return row_to_image


def _extract_rich_value_images(z, names, sheet_xml_path="xl/worksheets/sheet1.xml"):
    row_to_image = {}
    rels_path = "xl/richData/_rels/richValueRel.xml.rels"
    if rels_path not in names:
        return row_to_image
    rels_root = ET.fromstring(z.read(rels_path))
    rid_to_media = {}
    for rel in rels_root:
        rid = rel.get("Id", "")
        target = rel.get("Target", "")
        full = "xl/media/" + Path(target).name
        if full in names:
            rid_to_media[rid] = full

    rv_rel_root = ET.fromstring(z.read("xl/richData/richValueRel.xml"))
    NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    NS_RVR = "http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel"
    index_to_rid = []
    for rel in rv_rel_root:
        if rel.tag.endswith("}rel"):
            index_to_rid.append(rel.get(f"{{{NS_R}}}id", ""))

    rv_root = ET.fromstring(z.read("xl/richData/rdrichvalue.xml"))
    NS_RV = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"
    rv_index_to_media = {}
    for i, rv in enumerate(rv_root.findall(f"{{{NS_RV}}}rv")):
        values = [v.text for v in rv.findall(f"{{{NS_RV}}}v")]
        if not values:
            continue
        try:
            rel_idx = int(values[0])
        except (ValueError, TypeError):
            continue
        if rel_idx < len(index_to_rid):
            rid = index_to_rid[rel_idx]
            media = rid_to_media.get(rid, "")
            if media:
                rv_index_to_media[i] = media

    meta_root = ET.fromstring(z.read("xl/metadata.xml"))
    NS_META = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    NS_XLRD = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"
    future_to_rv = {}
    fut_meta = meta_root.find(f"{{{NS_META}}}futureMetadata[@name='XLRICHVALUE']")
    if fut_meta is not None:
        for fi, bk in enumerate(fut_meta.findall(f"{{{NS_META}}}bk")):
            rvb = bk.find(f".//{{{NS_XLRD}}}rvb")
            if rvb is not None:
                try:
                    future_to_rv[fi] = int(rvb.get("i", -1))
                except ValueError:
                    pass

    vm_to_future = {}
    val_meta = meta_root.find(f"{{{NS_META}}}valueMetadata")
    if val_meta is not None:
        for vi, bk in enumerate(val_meta.findall(f"{{{NS_META}}}bk")):
            for rc in bk.findall(f"{{{NS_META}}}rc"):
                try:
                    vm_to_future[vi] = int(rc.get("v", -1))
                except ValueError:
                    pass

    sheet_xml = z.read(sheet_xml_path).decode()
    for m in re.finditer(r'<c\s+r="([A-Z]+(\d+))"[^>]*\bvm="(\d+)"', sheet_xml):
        row_num = int(m.group(2))
        vm_idx = int(m.group(3)) - 1
        future_idx = vm_to_future.get(vm_idx, -1)
        rv_idx = future_to_rv.get(future_idx, -1)
        media_path = rv_index_to_media.get(rv_idx, "")
        if not media_path or row_num in row_to_image:
            continue
        suffix = Path(media_path).suffix.lower() or ".png"
        dest_name = f"row{row_num}{suffix}"
        dest_path = EMBED_DIR / dest_name
        with z.open(media_path) as src, open(dest_path, "wb") as dst:
            shutil.copyfileobj(src, dst)
        row_to_image[row_num] = dest_path
        print(f"  ✓ 画像抽出 (Excel 365): 行{row_num} → {dest_name}")
    return row_to_image


def _extract_drawing_images(z, names):
    row_to_image = {}
    drawing_candidates = sorted(n for n in names if re.match(r"xl/drawings/drawing\d+\.xml$", n))
    if not drawing_candidates:
        return row_to_image
    for drawing_xml_path in drawing_candidates:
        drawing_rels_path = re.sub(
            r"xl/drawings/(drawing\d+\.xml)$", r"xl/drawings/_rels/\1.rels", drawing_xml_path)
        rid_to_media = {}
        if drawing_rels_path in names:
            rels_root = ET.fromstring(z.read(drawing_rels_path))
            for rel in rels_root:
                rid = rel.get("Id", "")
                target = rel.get("Target", "")
                full = "xl/media/" + Path(target).name
                if full in names:
                    rid_to_media[rid] = full
        drawing_root = ET.fromstring(z.read(drawing_xml_path))
        NS = {
            "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
            "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
            "r":   "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        for anchor in drawing_root:
            local = anchor.tag.split("}")[-1]
            if local not in ("twoCellAnchor", "oneCellAnchor"):
                continue
            from_elem = anchor.find("xdr:from", NS)
            if from_elem is None:
                continue
            row_elem = from_elem.find("xdr:row", NS)
            if row_elem is None or not row_elem.text:
                continue
            excel_row = int(row_elem.text) + 1
            blip = anchor.find(".//a:blip", NS)
            if blip is None:
                continue
            rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", "")
            media_path = rid_to_media.get(rid, "")
            if not media_path or excel_row in row_to_image:
                continue
            suffix = Path(media_path).suffix.lower() or ".png"
            dest_name = f"row{excel_row}{suffix}"
            dest_path = EMBED_DIR / dest_name
            with z.open(media_path) as src, open(dest_path, "wb") as dst:
                shutil.copyfileobj(src, dst)
            row_to_image[excel_row] = dest_path
            print(f"  ✓ 画像抽出 (floating): 行{excel_row} → {dest_name}")
    return row_to_image


# ────────── アイコン画像を決める ──────────
def resolve_icon(domain: str, row_idx: int, embedded: dict):
    """icons/<ドメイン>.png があればそれを使う。なければ Excel の埋め込み画像を保存して使う。"""
    if not domain:
        return None
    for ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".ico"):
        if (ICONS_DIR / f"{domain}{ext}").exists():
            return f"icons/{domain}{ext}"
    src = embedded.get(row_idx)
    if src:
        ICONS_DIR.mkdir(exist_ok=True)
        dest = ICONS_DIR / f"{domain}{Path(src).suffix.lower() or '.png'}"
        shutil.copyfile(src, dest)
        print(f"  ✓ アイコン追加: 行{row_idx} → {dest.name}")
        return f"icons/{dest.name}"
    return None


# ────────── JS の CATS 配列を生成 ──────────
def generate_cats_js(categories_meta, category_rows, image_map):
    lines = ["const CATS = ["]
    fallback_idx = 0

    for cat, meta in categories_meta:
        rows = category_rows.get(cat, [])
        count = len(rows)
        icon  = meta["icon"]
        color = meta["color"]
        cat_esc = cat.replace("'", "\\'")

        lines.append(f"  {{icon:'{icon}', name:'{cat_esc}', color:'{color}', count:{count}, links:[")

        for row_idx, row in rows:
            def cell(v):
                if v is None: return ""
                s = str(v).strip()
                return "" if s.startswith("#") else s

            name   = cell(row[1]) if len(row) > 1 else ""
            url    = cell(row[2]) if len(row) > 2 else "#"
            url    = url or "#"
            desc   = cell(row[3]) if len(row) > 3 else ""
            domain = extract_domain(url)

            # アイコン画像
            img_path = resolve_icon(domain, row_idx, image_map)
            img_val  = f"'{img_path}'" if img_path else "null"

            name_esc   = name.replace("'", "\\'")
            url_esc    = url.replace("'", "\\'")
            desc_esc   = desc.replace("'", "\\'")
            domain_esc = domain.replace("'", "\\'")

            lines.append(
                f"    {{name:'{name_esc}', url:'{url_esc}', img:{img_val}, "
                f"desc:'{desc_esc}', domain:'{domain_esc}'}},"
            )

        lines.append("  ]},")

    lines.append("];")
    return "\n".join(lines)


# ────────── index.html を更新 ──────────
def update_html(cats_js: str):
    html = INDEX_HTML.read_text(encoding="utf-8")
    start_marker = "// ===== CATS START ====="
    end_marker   = "// ===== CATS END ====="
    si = html.find(start_marker)
    ei = html.find(end_marker)
    if si == -1 or ei == -1:
        print("エラー: index.html にマーカーが見つかりません。")
        print(f"  '{start_marker}' と '{end_marker}' が必要です。")
        sys.exit(1)
    new_html = (
        html[:si + len(start_marker)]
        + "\n"
        + cats_js
        + "\n"
        + html[ei:]
    )
    INDEX_HTML.write_text(new_html, encoding="utf-8")
    print("✅ index.html を更新しました。")


# ────────── メイン ──────────
def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"エラー: '{EXCEL_PATH}' が見つかりません。")
        sys.exit(1)

    print(f"📄 読み込み: {EXCEL_PATH}")
    DATA_SHEET = "リンク一覧"

    # Excel 埋め込み画像を抽出（icons/ に同じドメインの画像がない行だけ使う）
    print("🖼  Excel 埋め込み画像を確認中...")
    image_map = extract_embedded_images(EXCEL_PATH, DATA_SHEET)
    if not image_map:
        print("  (埋め込み画像なし)")


    wb = load_workbook(EXCEL_PATH, data_only=True)
    if DATA_SHEET in wb.sheetnames:
        ws = wb[DATA_SHEET]
        print(f"  シート「{DATA_SHEET}」を読み込みます")
    else:
        ws = wb.active
        print(f"  ⚠ 「{DATA_SHEET}」シートが見つからないため、アクティブシート「{ws.title}」を使います")

    ordered_cats = []
    category_rows = {}
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW
    ):
        if not any(row):
            continue
        cat = str(row[0]).strip() if row[0] else ""
        if not cat:
            continue
        if cat not in category_rows:
            ordered_cats.append(cat)
            category_rows[cat] = []
        category_rows[cat].append((row_idx, row))

    if not ordered_cats:
        print("エラー: カテゴリが見つかりません。A列にカテゴリ名を入力してください。")
        sys.exit(1)

    # カテゴリメタ情報を決定
    fallback_i = 0
    categories_meta = []
    for cat in ordered_cats:
        if cat in CATEGORY_META:
            meta = dict(CATEGORY_META[cat])
        else:
            meta = {
                "icon":  FALLBACK_ICONS[fallback_i % len(FALLBACK_ICONS)],
                "color": FALLBACK_COLORS[fallback_i % len(FALLBACK_COLORS)],
            }
            fallback_i += 1
            print(f"  新カテゴリ「{cat}」: アイコン {meta['icon']} を自動割り当て")
        categories_meta.append((cat, meta))

    cats_js = generate_cats_js(categories_meta, category_rows, image_map)
    update_html(cats_js)

    print(f"\n📊 集計:")
    total = 0
    for cat, _ in categories_meta:
        n = len(category_rows.get(cat, []))
        total += n
        print(f"  {cat}: {n} 件")
    print(f"  合計: {total} 件")


if __name__ == "__main__":
    main()
